from datetime import date
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from httpx2 import Response
from openpyxl import Workbook, load_workbook
from sqlalchemy import Engine
from sqlalchemy.orm import Session

from venergy_api.database import Base
from venergy_api.models import DimUc, OrdemEvento
from venergy_api.planilha_campo import COLUNAS_EXPORTACAO

ROTA = "/ordens-inspecao/planilha"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CABECALHO = [
    "id_inspecao",
    "id_uc",
    "data",
    "inspetor",
    "resultado",
    "tipo_irregularidade",
    "kwh_recuperado",
]
UC_BASE = {
    "municipio": "Recife",
    "bairro": "Boa Viagem",
    "classe": "B1",
    "classe_cadastro": "Residencial",
    "cod_alimentador": "AL-01",
    "possui_gd": False,
    "consumo_medio_kwh_12m": 150.0,
    "data_ligacao": date(2018, 6, 1),
}


@pytest.fixture(autouse=True)
def stored_ucs(engine: Engine) -> None:
    Base.metadata.create_all(engine)
    with Session(engine) as session, session.begin():
        session.add_all([DimUc(cod_uc=cod, **UC_BASE) for cod in ("0001", "0002", "0003", "0004")])


def confirmada(cod_uc: str, inspecao: str = "INS-1") -> list:
    return [inspecao, cod_uc, "2025-08-12", "Fulano", "CONFIRMADA", "DESVIO DIRETO", 120.5]


def regular(cod_uc: str, inspecao: str = "INS-2") -> list:
    return [inspecao, cod_uc, "2025-08-13", "Ciclana", "SEM IRREGULARIDADE", None, 0]


def xlsx(*linhas: list) -> bytes:
    livro = Workbook()
    aba = livro.active
    aba.title = "ago-set"
    aba.append(CABECALHO)
    for linha in linhas:
        aba.append(linha)
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def enviar(client: TestClient, conteudo: bytes, nome: str = "campo.xlsx") -> Response:
    return client.post(ROTA, files={"arquivo": (nome, conteudo, XLSX)})


def gerar_ordem(client: TestClient, cod_uc: str, *, em_campo: bool = True) -> int:
    ordem_id = client.post("/ordens-inspecao", json={"cod_ucs": [cod_uc]}).json()[0]["id"]
    if em_campo:
        assert client.patch(f"/ordens-inspecao/{ordem_id}", json={"status": "EM_CAMPO"}).is_success
    return ordem_id


def detalhe(client: TestClient, ordem_id: int) -> dict:
    return client.get(f"/ordens-inspecao/{ordem_id}").json()


def test_concludes_the_order_that_is_in_the_field(client: TestClient) -> None:
    ordem_id = gerar_ordem(client, "0001")

    response = enviar(client, xlsx(confirmada("0001")))

    assert response.status_code == 200, response.text
    assert response.json() == {
        "linhas_lidas": 1,
        "concluidas": [
            {
                "aba": "ago-set",
                "linha": 2,
                "cod_uc": "0001",
                "ordem_id": ordem_id,
                "resultado": "IRREGULAR_CONFIRMADA",
            }
        ],
        "rejeitadas": [],
    }
    ordem = detalhe(client, ordem_id)
    assert (ordem["status"], ordem["resultado"]) == ("CONCLUIDA", "IRREGULAR_CONFIRMADA")
    assert ordem["observacao"].startswith("Planilha de campo: desvio direto, 120,5 kWh")


def test_the_history_records_who_uploaded_the_file(client: TestClient) -> None:
    ordem_id = gerar_ordem(client, "0001")

    enviar(client, xlsx(regular("0001")))

    ultimo = detalhe(client, ordem_id)["eventos"][-1]
    assert (ultimo["status_anterior"], ultimo["status_novo"]) == ("EM_CAMPO", "CONCLUIDA")
    assert ultimo["autor_login"] == "teste"
    assert "sem irregularidade" in ultimo["observacao"]


def test_a_generated_order_is_rejected_and_left_untouched(client: TestClient) -> None:
    ordem_id = gerar_ordem(client, "0001", em_campo=False)

    corpo = enviar(client, xlsx(confirmada("0001"))).json()

    assert corpo["concluidas"] == []
    assert corpo["rejeitadas"][0]["cod_uc"] == "0001"
    assert "não foi enviada a campo" in corpo["rejeitadas"][0]["motivo"]
    assert detalhe(client, ordem_id)["status"] == "GERADA"


def test_a_uc_without_any_order_is_rejected(client: TestClient) -> None:
    corpo = enviar(client, xlsx(confirmada("0002"))).json()

    assert corpo["concluidas"] == []
    assert "Nenhuma ordem" in corpo["rejeitadas"][0]["motivo"]


def test_a_uc_that_does_not_exist_is_rejected(client: TestClient) -> None:
    corpo = enviar(client, xlsx(confirmada("9999"))).json()

    assert corpo["concluidas"] == []
    assert corpo["rejeitadas"][0]["cod_uc"] == "9999"


def test_a_cancelled_order_is_rejected(client: TestClient) -> None:
    ordem_id = gerar_ordem(client, "0001", em_campo=False)
    client.patch(f"/ordens-inspecao/{ordem_id}", json={"status": "CANCELADA"})

    corpo = enviar(client, xlsx(confirmada("0001"))).json()

    assert "cancelada" in corpo["rejeitadas"][0]["motivo"]


def test_sending_the_same_file_twice_changes_nothing_the_second_time(
    client: TestClient, engine: Engine
) -> None:
    ordem_id = gerar_ordem(client, "0001")
    arquivo = xlsx(confirmada("0001"))
    enviar(client, arquivo)
    with Session(engine) as session:
        eventos_antes = session.query(OrdemEvento).count()

    corpo = enviar(client, arquivo).json()

    assert corpo["concluidas"] == []
    assert "já concluída" in corpo["rejeitadas"][0]["motivo"]
    with Session(engine) as session:
        assert session.query(OrdemEvento).count() == eventos_antes
    assert detalhe(client, ordem_id)["status"] == "CONCLUIDA"


def test_valid_rows_are_applied_even_when_others_are_rejected(client: TestClient) -> None:
    primeira = gerar_ordem(client, "0001")
    gerada = gerar_ordem(client, "0002", em_campo=False)
    terceira = gerar_ordem(client, "0003")
    invalida = ["INS-9", "0004", None, None, "FRAUDE", None, 0]

    corpo = enviar(
        client, xlsx(confirmada("0001"), regular("0002"), regular("0003"), invalida)
    ).json()

    assert corpo["linhas_lidas"] == 4
    assert [c["ordem_id"] for c in corpo["concluidas"]] == [primeira, terceira]
    assert sorted(r["cod_uc"] for r in corpo["rejeitadas"]) == ["0002", "0004"]
    assert detalhe(client, gerada)["status"] == "GERADA"
    assert detalhe(client, terceira)["resultado"] == "REGULAR"


@pytest.mark.parametrize(
    ("conteudo", "trecho"),
    [(b"nao e uma planilha", "XLSX"), (xlsx(), "vazia")],
)
def test_an_unusable_file_returns_422_and_records_nothing(
    client: TestClient, conteudo: bytes, trecho: str
) -> None:
    ordem_id = gerar_ordem(client, "0001")

    response = enviar(client, conteudo)

    assert response.status_code == 422
    assert trecho in response.json()["detail"]
    assert detalhe(client, ordem_id)["status"] == "EM_CAMPO"


def test_a_request_without_the_file_returns_422(client: TestClient) -> None:
    assert client.post(ROTA).status_code == 422


def baixar(client: TestClient) -> Workbook:
    response = client.get(ROTA)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    assert "ordens-em-campo.xlsx" in response.headers["content-disposition"]
    return load_workbook(BytesIO(response.content))


def test_export_lists_only_the_orders_in_the_field_with_the_upload_header(
    client: TestClient,
) -> None:
    em_campo = gerar_ordem(client, "0001")
    gerar_ordem(client, "0002", em_campo=False)
    concluida = gerar_ordem(client, "0003")
    client.patch(f"/ordens-inspecao/{concluida}", json={"status": "CANCELADA"})

    linhas = list(baixar(client).active.iter_rows(values_only=True))

    assert linhas[0] == tuple(COLUNAS_EXPORTACAO)
    assert linhas[1:] == [
        (f"OI-{em_campo}", "0001", None, None, None, None, None, "Recife", "Boa Viagem")
    ]


def test_export_without_orders_in_the_field_has_only_the_header(client: TestClient) -> None:
    gerar_ordem(client, "0001", em_campo=False)

    assert len(list(baixar(client).active.iter_rows())) == 1


def test_the_exported_file_filled_in_comes_back_through_the_upload(client: TestClient) -> None:
    primeira = gerar_ordem(client, "0001")
    segunda = gerar_ordem(client, "0002")
    livro = baixar(client)
    aba = livro.active
    aba["E2"], aba["F2"], aba["G2"] = "CONFIRMADA", "BY-PASS", "80,5"
    aba["E3"], aba["G3"] = "SEM IRREGULARIDADE", 0
    arquivo = BytesIO()
    livro.save(arquivo)

    corpo = enviar(client, arquivo.getvalue()).json()

    assert [c["ordem_id"] for c in corpo["concluidas"]] == [primeira, segunda]
    assert corpo["rejeitadas"] == []
    ordem = detalhe(client, primeira)
    assert ordem["resultado"] == "IRREGULAR_CONFIRMADA"
    assert "(inspeção OI-" in ordem["observacao"]


def test_the_export_restricts_the_columns_the_field_team_fills(client: TestClient) -> None:
    gerar_ordem(client, "0001")
    aba = baixar(client).active

    por_coluna = {str(v.sqref).split(":")[0][0]: v for v in aba.data_validations.dataValidation}

    assert por_coluna["E"].type == "list"
    assert por_coluna["E"].formula1 == '"CONFIRMADA,SEM IRREGULARIDADE"'
    assert por_coluna["C"].type == "date"
    assert por_coluna["G"].type == "decimal"


def _salvar(livro: Workbook) -> bytes:
    arquivo = BytesIO()
    livro.save(arquivo)
    return arquivo.getvalue()


def test_an_old_sheet_does_not_conclude_a_newer_order_of_the_same_uc(client: TestClient) -> None:
    antiga = gerar_ordem(client, "0001")
    client.patch(f"/ordens-inspecao/{antiga}", json={"status": "CANCELADA"})
    nova = gerar_ordem(client, "0001")
    livro = Workbook()
    livro.active.append(["id_inspecao", "id_uc", "resultado"])
    livro.active.append([f"OI-{antiga}", "0001", "SEM IRREGULARIDADE"])
    saida = BytesIO()
    livro.save(saida)

    corpo = enviar(client, saida.getvalue()).json()

    assert corpo["concluidas"] == []
    assert f"ordem {antiga}" in corpo["rejeitadas"][0]["motivo"]
    assert detalhe(client, nova)["status"] == "EM_CAMPO"
