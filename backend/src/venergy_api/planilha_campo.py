"""Leitura da planilha que a equipe de campo preenche à mão (formato de `inspecoes.xlsx`).

Fica separada do banco e da API para ser testada só com planilhas montadas em memória. As regras
de limpeza repetem as do pipeline (docs/qualidade-dados.md, INSP-01 a INSP-07): a planilha que
chega pelo portal precisa ser tratada como a que passa pela Silver, senão o mesmo resultado
seria lido de dois jeitos.
"""

import re
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

from venergy_api.ordens import ResultadoOrdem

MAX_BYTES = 5 * 1024 * 1024
# O arquivo do case tem cerca de 1.400 linhas no ano todo; o teto só barra arquivo errado.
MAX_LINHAS = 5000
# Mesmo limite da observação digitada no modal (AtualizarOrdemSchema).
MAX_OBSERVACAO = 1000

COLUNAS_OBRIGATORIAS = ("id_uc", "resultado")

RESULTADOS = {
    "IRREGULARIDADE CONFIRMADA": ResultadoOrdem.IRREGULAR_CONFIRMADA,
    "CONFIRMADA": ResultadoOrdem.IRREGULAR_CONFIRMADA,
    "SEM IRREGULARIDADE": ResultadoOrdem.REGULAR,
    "NADA CONSTA": ResultadoOrdem.REGULAR,
}

TIPO_AUSENTE = {"", "-", "N/A", "NULL"}

FORMATOS_DE_DATA = (
    (re.compile(r"\d{4}-\d{2}-\d{2}"), "%Y-%m-%d"),
    (re.compile(r"\d{2}\.\d{2}\.\d{4}"), "%d.%m.%Y"),
    (re.compile(r"\d{2}/\d{2}/\d{4}"), "%d/%m/%Y"),
    (re.compile(r"\d{2}-\d{2}-\d{4}"), "%m-%d-%Y"),
)


class PlanilhaInvalida(Exception):
    """O arquivo inteiro não serve: nada é gravado e a pessoa corrige e envia de novo."""


@dataclass(frozen=True)
class LinhaCampo:
    aba: str
    linha: int
    cod_uc: str
    resultado: ResultadoOrdem
    observacao: str


@dataclass(frozen=True)
class LinhaRejeitada:
    aba: str
    linha: int
    cod_uc: str | None
    motivo: str


@dataclass(frozen=True)
class PlanilhaLida:
    linhas_lidas: int
    validas: list[LinhaCampo]
    rejeitadas: list[LinhaRejeitada]


class _LinhaInvalida(Exception):
    pass


def ler_planilha(conteudo: bytes) -> PlanilhaLida:
    if len(conteudo) > MAX_BYTES:
        raise PlanilhaInvalida(f"Arquivo grande demais: o limite é {MAX_BYTES // (1024 * 1024)} MB")
    try:
        livro = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as error:
        raise PlanilhaInvalida("O arquivo não é uma planilha XLSX válida") from error

    linhas_lidas = 0
    validas: list[LinhaCampo] = []
    rejeitadas: list[LinhaRejeitada] = []
    ucs_vistas: set[str] = set()
    try:
        for aba in livro.worksheets:
            linhas = enumerate(aba.iter_rows(values_only=True), start=1)
            colunas = _achar_colunas(aba.title, linhas)
            if colunas is None:
                continue
            for numero, valores in linhas:
                if all(_vazio(valor) for valor in valores):
                    continue
                linhas_lidas += 1
                if linhas_lidas > MAX_LINHAS:
                    raise PlanilhaInvalida(f"Planilha com mais de {MAX_LINHAS} linhas")
                celula = {nome: _valor(valores, posicao) for nome, posicao in colunas.items()}
                cod_uc = _texto(celula.get("id_uc"))
                try:
                    linha = _ler_linha(aba.title, numero, cod_uc, celula)
                    if linha.cod_uc in ucs_vistas:
                        raise _LinhaInvalida("UC repetida na planilha: vale a primeira linha")
                except _LinhaInvalida as erro:
                    rejeitadas.append(LinhaRejeitada(aba.title, numero, cod_uc or None, str(erro)))
                    continue
                ucs_vistas.add(linha.cod_uc)
                validas.append(linha)
    finally:
        livro.close()

    if linhas_lidas == 0:
        raise PlanilhaInvalida("Planilha vazia: nenhuma linha de inspeção encontrada")
    return PlanilhaLida(linhas_lidas, validas, rejeitadas)


def _achar_colunas(aba: str, linhas) -> dict[str, int] | None:
    for _, valores in linhas:
        if all(_vazio(valor) for valor in valores):
            continue
        nomes = [
            re.sub(r"\s+", "_", str(valor).strip().lower()) if valor else "" for valor in valores
        ]
        faltando = [coluna for coluna in COLUNAS_OBRIGATORIAS if coluna not in nomes]
        if faltando:
            raise PlanilhaInvalida(f"Aba {aba}: faltam as colunas {', '.join(faltando)}")
        return {nome: posicao for posicao, nome in enumerate(nomes) if nome}
    return None


def _ler_linha(aba: str, numero: int, cod_uc: str, celula: dict[str, object]) -> LinhaCampo:
    if not cod_uc:
        raise _LinhaInvalida("UC em branco")
    bruto = " ".join(_texto(celula.get("resultado")).upper().split())
    if not bruto:
        raise _LinhaInvalida("Resultado em branco")
    resultado = RESULTADOS.get(bruto)
    if resultado is None:
        raise _LinhaInvalida(f"Resultado não reconhecido: {bruto.lower()}")

    tipo = _tipo(celula.get("tipo_irregularidade"))
    kwh = _kwh(celula.get("kwh_recuperado"))
    data = _data(celula.get("data"))
    inspetor = " ".join(_texto(celula.get("inspetor")).split())
    inspecao = _texto(celula.get("id_inspecao"))

    if resultado is ResultadoOrdem.IRREGULAR_CONFIRMADA:
        if tipo is None:
            raise _LinhaInvalida("Irregularidade confirmada sem o tipo de irregularidade")
        if kwh is None or kwh <= 0:
            raise _LinhaInvalida(
                "Irregularidade confirmada precisa de kWh recuperado maior que zero"
            )
    else:
        if tipo is not None:
            raise _LinhaInvalida("Sem irregularidade com tipo de irregularidade preenchido")
        if kwh is not None and kwh > 0:
            raise _LinhaInvalida("Sem irregularidade com kWh recuperado maior que zero")

    partes = ["sem irregularidade" if tipo is None else tipo.lower()]
    if kwh is not None and tipo is not None:
        partes.append(f"{_formatar_kwh(kwh)} kWh recuperados")
    if inspetor:
        partes.append(f"inspetor {inspetor}")
    if data is not None:
        partes.append(data.strftime("%d/%m/%Y"))
    observacao = "Planilha de campo: " + ", ".join(partes)
    if inspecao:
        observacao += f" (inspeção {inspecao})"
    return LinhaCampo(aba, numero, cod_uc, resultado, observacao[:MAX_OBSERVACAO])


def _vazio(valor: object) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())


def _valor(valores: tuple, posicao: int) -> object:
    return valores[posicao] if posicao < len(valores) else None


def _texto(valor: object) -> str:
    if valor is None:
        return ""
    # O Excel guarda número como float; "17.0" não é código de UC.
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _tipo(valor: object) -> str | None:
    texto = _texto(valor)
    return None if texto.upper() in TIPO_AUSENTE else texto


def _kwh(valor: object) -> Decimal | None:
    texto = _texto(valor)
    if not texto:
        return None
    # "1.234,5": com vírgula, o ponto é separador de milhar (INSP-07).
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        numero = Decimal(texto)
    except InvalidOperation:
        raise _LinhaInvalida(f"kWh recuperado inválido: {_texto(valor)}") from None
    if not numero.is_finite():
        raise _LinhaInvalida(f"kWh recuperado inválido: {_texto(valor)}")
    return numero


def _data(valor: object) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = _texto(valor)
    if not texto:
        return None
    for padrao, formato in FORMATOS_DE_DATA:
        if padrao.fullmatch(texto):
            try:
                return datetime.strptime(texto, formato).date()
            except ValueError:
                break
    raise _LinhaInvalida(f"Data inválida: {texto}")


def _formatar_kwh(kwh: Decimal) -> str:
    texto = f"{kwh:,.3f}".rstrip("0").rstrip(".")
    return texto.replace(",", "_").replace(".", ",").replace("_", ".")


TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# As sete primeiras são as do formato de inspecoes.xlsx: a equipe preenche e o upload lê. Município
# e bairro vão ao final só para quem está em campo saber onde fica a unidade; o leitor as ignora.
COLUNAS_EXPORTACAO = (
    "id_inspecao",
    "id_uc",
    "data",
    "inspetor",
    "resultado",
    "tipo_irregularidade",
    "kwh_recuperado",
    "municipio",
    "bairro",
)


# Linhas que a validação das colunas cobre: sobra folga para a equipe acrescentar linhas.
LINHAS_VALIDADAS = 2000
OPCOES_RESULTADO = ("CONFIRMADA", "SEM IRREGULARIDADE")


def _validar_colunas(aba) -> None:
    """Seleção no lugar de digitação: o Excel barra o que o upload rejeitaria."""
    fim = LINHAS_VALIDADAS
    resultado = DataValidation(
        type="list",
        formula1='"' + ",".join(OPCOES_RESULTADO) + '"',
        allow_blank=True,
        errorTitle="Resultado inválido",
        error="Escolha CONFIRMADA ou SEM IRREGULARIDADE.",
        promptTitle="Resultado",
        prompt="CONFIRMADA: achou irregularidade. SEM IRREGULARIDADE: não achou.",
    )
    data = DataValidation(
        type="date",
        operator="greaterThan",
        formula1="36526",
        allow_blank=True,
        errorTitle="Data inválida",
        error="Informe uma data, como 12/08/2025.",
    )
    kwh = DataValidation(
        type="decimal",
        operator="greaterThanOrEqual",
        formula1="0",
        allow_blank=True,
        errorTitle="kWh inválido",
        error="Informe um número maior ou igual a zero.",
    )
    for validacao, coluna in ((data, "C"), (resultado, "E"), (kwh, "G")):
        validacao.showErrorMessage = True
        validacao.showInputMessage = validacao is resultado
        aba.add_data_validation(validacao)
        validacao.add(f"{coluna}2:{coluna}{fim}")
    aba.column_dimensions["C"].number_format = "dd/mm/yyyy"


def montar_planilha(ordens: Iterable[tuple[int, str, str | None, str | None]]) -> bytes:
    """Planilha das ordens em campo, no formato que o upload lê de volta.

    Cada item é (id da ordem, cod_uc, município, bairro). O código da inspeção sai como
    `OI-<id>`, para a observação da ordem concluída apontar de volta para ela.
    """
    livro = Workbook()
    aba = livro.active
    aba.title = "em-campo"
    aba.append(list(COLUNAS_EXPORTACAO))
    for ordem_id, cod_uc, municipio, bairro in ordens:
        aba.append([f"OI-{ordem_id}", cod_uc, None, None, None, None, None, municipio, bairro])
    _validar_colunas(aba)
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()
